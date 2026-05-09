# Author: Sergio Flores - Viti Science - VS-2026-PFI-001
"""Inspect master_workbook_v28 sheet list, region naming convention, and any
existing 'Methods notes' or AVA-related sheet — so the v29 insert plays nicely
with the existing structure."""

from pathlib import Path
import sys
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

WB = Path(r"C:/Users/sergi/OneDrive/Desktop/Viti Science/pfi-se-regional-mapping-main/delivery/06_master_workbook_v28.xlsx")

wb = openpyxl.load_workbook(WB, read_only=True, data_only=True)
print(f"v28 sheets ({len(wb.sheetnames)}):")
for i, name in enumerate(wb.sheetnames, 1):
    ws = wb[name]
    print(f"  {i:>2}. {name}  ({ws.max_row}x{ws.max_column})")

# Look for Methods notes specifically
candidates = [s for s in wb.sheetnames if "method" in s.lower() or "notes" in s.lower()
              or "ava" in s.lower() or "build" in s.lower()]
print(f"\nMethods/notes/AVA candidates: {candidates}")

# Peek at one likely region-keyed sheet to confirm region naming convention
region_sheets = [s for s in wb.sheetnames if "region" in s.lower() or "stockholm" in s.lower()][:3]
for s in region_sheets:
    ws = wb[s]
    print(f"\n--- Sample from sheet '{s}' (first 5 rows) ---")
    for row in ws.iter_rows(min_row=1, max_row=5, values_only=True):
        print("   ", row)

wb.close()
