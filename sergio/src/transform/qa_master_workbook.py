# Author: Sergio Flores - Viti Science - VS-2026-PFI-001
"""Self-check the master workbook: validate sheet contents, spot anomalies, sanity-check totals."""
import sys
sys.stdout.reconfigure(encoding="utf-8")
import openpyxl
from pathlib import Path

WB = Path("C:/Users/sergi/OneDrive/Desktop/Viti Science/pfi-se-regional-mapping-main/working/data/master/master_workbook_v01.xlsx")

wb = openpyxl.load_workbook(WB, data_only=True)
print(f"Workbook: {WB.name}")
print(f"Sheets: {wb.sheetnames}")
print()

issues = []

def check(label, cond, detail=""):
    status = "OK" if cond else "FAIL"
    print(f"  [{status}] {label}{(' — ' + detail) if detail else ''}")
    if not cond:
        issues.append(f"{label}{(' — ' + detail) if detail else ''}")

# === Region master ===
print("=== Region master sheet ===")
ws = wb["Region master"]
rows = list(ws.iter_rows(values_only=True))
hdr = rows[0]
data = rows[1:]
check("21 region rows", len(data) == 21, f"got {len(data)}")
check("Headers present", hdr[0] == "Region code")

# Check totals
total_pop = sum(r[3] for r in data if isinstance(r[3], (int,float)))
check("National pop ≈ 10.5–10.7M", 10_500_000 <= total_pop <= 10_700_000, f"sum = {total_pop:,}")

total_brp = sum(r[9] for r in data if isinstance(r[9], (int,float)))  # BRP total mnkr
check("National BRP ≈ 5–7 trillion SEK", 5_000_000 <= total_brp <= 7_000_000, f"sum = {total_brp:,.0f} mnkr")

# Stockholm should be ~2.4M, Gotland ~60k
sthlm = next(r for r in data if r[1] == "Region Stockholm")
gotland = next(r for r in data if r[1] == "Region Gotland")
check("Stockholm pop ~2.4M", 2_400_000 <= sthlm[3] <= 2_500_000, f"got {sthlm[3]:,}")
check("Gotland pop ~60k", 55_000 <= gotland[3] <= 65_000, f"got {gotland[3]:,}")

# Healthcare cost / inv should be 30k-45k SEK roughly
hc_costs = [r[11] for r in data if isinstance(r[11], (int,float))]
check("Healthcare cost/inv all in 25–55k SEK range",
      all(25_000 <= c <= 55_000 for c in hc_costs),
      f"min={min(hc_costs):,.0f}, max={max(hc_costs):,.0f}")

# Pharma förmån ≤ Pharma total (logically)
forman_le_total = []
for r in data:
    if isinstance(r[14], (int,float)) and isinstance(r[15], (int,float)):
        forman_le_total.append((r[1], r[14], r[15], r[14] <= r[15]))
n_violations = sum(1 for x in forman_le_total if not x[3])
check("Förmån ≤ Total pharma for all regions", n_violations == 0,
      f"{n_violations} violations")
# But also confirm they're equal (the anomaly we flagged)
n_equal = sum(1 for x in forman_le_total if abs(x[1] - x[2]) < 1)
check("Förmån vs Total pharma — anomaly check", n_equal < 21,
      f"{n_equal}/21 regions have förmån = total (suggests no rekvisition split)")
print(f"     This is the methodology issue flagged in findings_kolada_indicators.md")

# === Per-capita view ===
print("\n=== Per-capita view sheet ===")
ws = wb["Per-capita view"]
rows = list(ws.iter_rows(values_only=True))
data = rows[1:]
check("21 region rows", len(data) == 21, f"got {len(data)}")

# Vaccine sell-in / inv per year — should be in tens-hundreds of SEK
vacc_per_yr = [r[8] for r in data if isinstance(r[8], (int,float))]
check("Vaccine sell-in/inv/year in 10–200 SEK range",
      all(10 <= v <= 200 for v in vacc_per_yr),
      f"min={min(vacc_per_yr):.1f}, max={max(vacc_per_yr):.1f}")

# Top region
ranked = sorted(data, key=lambda r: -(r[8] or 0))
print(f"     Top 3 vaccine spend/inv/yr: " +
      ", ".join(f"{r[1]} ({r[8]:.1f} SEK)" for r in ranked[:3]))
print(f"     Bottom 3:                    " +
      ", ".join(f"{r[1]} ({r[8]:.1f} SEK)" for r in ranked[-3:]))

# === Vaccine market share ===
print("\n=== Vaccine market share sheet ===")
ws = wb["Vaccine market share"]
rows = list(ws.iter_rows(values_only=True))
data = rows[1:]
check("21 region rows", len(data) == 21)

# Pfizer share % should be 0–100
shares = []
for r in data:
    for col in [4, 7, 10]:  # share % columns
        if isinstance(r[col], (int,float)):
            shares.append(r[col])
check("All Pfizer share % in 0–100", all(0 <= s <= 100 for s in shares),
      f"min={min(shares):.1f}, max={max(shares):.1f}")

# Show range per ATC class
print(f"     RSV (Abrysvo vs Arexvy):")
print(f"       Pfizer share range: {min(r[4] for r in data if isinstance(r[4],(int,float))):.1f}% – {max(r[4] for r in data if isinstance(r[4],(int,float))):.1f}%")
print(f"     Pneumococcal (Prevenar 13/20 vs MSD's Vaxneuvance/Capvaxive):")
print(f"       Pfizer share range: {min(r[7] for r in data if isinstance(r[7],(int,float))):.1f}% – {max(r[7] for r in data if isinstance(r[7],(int,float))):.1f}%")
print(f"     TBE (FSME-IMMUN vs Encepur):")
print(f"       Pfizer share range: {min(r[10] for r in data if isinstance(r[10],(int,float))):.1f}% – {max(r[10] for r in data if isinstance(r[10],(int,float))):.1f}%")

# === Vaccine sales detail ===
print("\n=== Vaccine sales detail sheet ===")
ws = wb["Vaccine sales detail"]
rows = list(ws.iter_rows(values_only=True))
data = rows[1:]
# Should be ~21 regions × 10 products = ~210, but some products won't have data in all regions
# Plus 1 "Unknown county" row
check("Detail rows present", 50 < len(data) < 250, f"got {len(data)} rows")

# Total sell-in across all = should match IQVIA file total (~2.5B SEK)
# Note: includes "Unknown county" row
total = sum(r[5] for r in data if isinstance(r[5], (int,float)))
check("Total sell-in ≈ 2.5B SEK", 2_400_000_000 <= total <= 2_600_000_000,
      f"sum = {total:,.0f} SEK")

# === Data sources sheet ===
print("\n=== Data sources sheet ===")
ws = wb["Data sources"]
rows = list(ws.iter_rows(values_only=True))
check("Sources documented", len(rows) > 5, f"{len(rows)-1} source rows")

# === Summary ===
print(f"\n{'='*60}")
print(f"QA SUMMARY: {len(issues)} issues found")
if issues:
    for i in issues:
        print(f"  - {i}")
else:
    print("All checks passed. Workbook ready for review.")
print(f"{'='*60}")

# === Next-step suggestions ===
print("\nObservations / next-step opportunities:")
print("- Förmån/rekvisition split is identical for all (or most) regions in Kolada → confirms need to source elsewhere (SKR or annual reports)")
print("- Vaccine market share by ATC is interesting per-region signal — competitive insight is real")
print("- BRP per inv variation (Stockholm 833 vs Gävleborg 434) supports the regional-economic-context narrative")
print("- Healthcare cost/inv variation is narrower than expected — small range likely reflects national cost-equalisation system")
