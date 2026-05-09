# Author: Sergio Flores - Viti Science - VS-2026-PFI-001
"""Rebuild deck_model.json with ALL fields needed for the appendix slides."""
import openpyxl, json, sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('delivery/06_master_workbook_v31.xlsx', read_only=True, data_only=True)

ws_rm = wb['Region master']
hdr_rm = [c.value for c in next(ws_rm.iter_rows())]
rm = {}
for row in ws_rm.iter_rows(min_row=2, values_only=True):
    if row[1] and ('Region' in str(row[1]) or 'Vastra' in str(row[1]).replace('ä', 'a')):
        rm[row[1]] = dict(zip(hdr_rm, row))

ws_pp = wb['Population projections']
hdr_pp = [c.value for c in next(ws_pp.iter_rows())]
pp = {}
for row in ws_pp.iter_rows(min_row=2, values_only=True):
    if row[1] and ('Region' in str(row[1]) or 'Vastra' in str(row[1]).replace('ä', 'a')):
        pp[row[1]] = dict(zip(hdr_pp, row))

ws_pf = wb['Pfizer total footprint']
hdr_pf = [c.value for c in next(ws_pf.iter_rows())]
pf = {}
for row in ws_pf.iter_rows(min_row=2, values_only=True):
    if row[0] and ('Region' in str(row[0]) or 'Vastra' in str(row[0]).replace('ä', 'a')):
        pf[row[0]] = dict(zip(hdr_pf, row))

ws_vac = wb['Vaccine market share']
vac = {}
for row in ws_vac.iter_rows(min_row=2, values_only=True):
    if row[1] and ('Region' in str(row[1]) or 'Vastra' in str(row[1]).replace('ä', 'a')):
        vac[row[1]] = {'rsv_share': row[4], 'pneu_share': row[7], 'tbe_share': row[10]}

ws_eq = wb['Equity composite']
hdr_eq = [c.value for c in next(ws_eq.iter_rows())]
eq = {}
for row in ws_eq.iter_rows(min_row=2, values_only=True):
    if row[0] and ('Region' in str(row[0]) or 'Vastra' in str(row[0]).replace('ä', 'a')):
        eq[row[0]] = dict(zip(hdr_eq, row))

ws_tp = wb['Top product per region']
hdr_tp = ['Region', 'Healthcare region', 'Top Pfizer product', 'Top Score', '2nd priority', '2nd Score', '3rd priority', '3rd Score']
tp = {}
for row in ws_tp.iter_rows(min_row=2, values_only=True):
    if row[0] and ('Region' in str(row[0]) or 'Vastra' in str(row[0]).replace('ä', 'a')):
        tp[row[0]] = dict(zip(hdr_tp, row))

ws_ar = wb['Regional archetypes']
hdr_ar = ['Region', 'Healthcare region', 'Primary archetype', 'Secondary archetypes', 'All tags', 'Pop 65+ %', '65+ growth 2024 to 2040 %', 'Median wait spec', 'Vyndaqel /100k', 'Confidence']
ar = {}
for row in ws_ar.iter_rows(min_row=5, values_only=True):
    if row[0] and ('Region' in str(row[0]) or 'Vastra' in str(row[0]).replace('ä', 'a')):
        ar[row[0]] = dict(zip(hdr_ar, row))

# Build full deck model
rows = []
for region in rm:
    r = rm[region]
    p = pp.get(region, {})
    f = pf.get(region, {})
    v = vac.get(region, {})
    e = eq.get(region, {})
    t = tp.get(region, {})
    a = ar.get(region, {})

    population = r.get('Population') or 0
    pop_65 = r.get('Pop 65+') or 0
    hc_cost = r.get('Healthcare cost per capita (SEK)')
    pharma = r.get('Pharma total cost per capita (SEK)')
    breast_rate = r.get('Breast ca rate /100k')
    prostate_rate = r.get('Prostate ca rate /100k')
    lung_rate = r.get('Lung ca incidence /100k')

    # Find the 65+ growth column dynamically (Swedish chars in the header)
    growth_key = None
    for k in p.keys():
        if k and 'growth' in str(k).lower():
            growth_key = k
            break

    rows.append({
        'region': region,
        'short': region.replace('Region ', '').replace('Västra Götalandsregionen', 'VGR'),
        'svr': r.get('Healthcare region'),
        # Demographics
        'pop': population,
        'pop_65_share': round(pop_65 / population * 100, 1) if population else None,
        'growth_65_2040': p.get(growth_key) if growth_key else None,
        'grp_per_capita_ksek': r.get('GRP per capita (kSEK)'),
        'foreign_born_pct': r.get('Foreign-born % (2024)'),
        'post_secondary_pct': None,  # Will set from a different source if needed
        # Population Projections
        'pop_2040': p.get('Pop 2040'),
        'pop_65_2040': p.get('65+ 2040'),
        'pop_2050': p.get('Pop 2050'),
        'births_2024': r.get('Births 2024'),
        'pop_women_15_44': r.get('Pop women 15-44'),
        'pop_0_1': r.get('Pop 0-1'),
        # Healthcare Budget
        'hc_cost_per_capita': hc_cost,
        'pharma_per_capita': pharma,
        'primary_care_per_capita': r.get('Primary care cost per capita (SEK)'),
        'pharma_pct_hc': round(pharma / hc_cost * 100, 1) if hc_cost and pharma else None,
        'vaccine_sellin_per_capita': r.get('Vaccine sell-in per capita (SEK/year)'),
        'specialists': r.get('Specialist physicians (#)'),
        # Health Equity Indicators
        'premature_mortality_25_64': None,  # Set below
        'hc_amenable_mortality': None,
        'suicide_25plus': None,
        'daily_smokers_pct': r.get('Daily smokers %'),
        'overweight_obese_pct': r.get('Overweight/obese %'),
        'equity_rank': e.get('SES composite rank (1=highest need)'),
        # Cancer burden
        'breast_ca_rate': breast_rate,
        'prostate_ca_rate': prostate_rate,
        'lung_ca_incidence': lung_rate,
        'lung_ca_mortality': r.get('Lung ca mortality /100k'),
        'breast_implied_cases': breast_rate * population / 100000 if breast_rate else None,
        'prostate_implied_cases': prostate_rate * population / 100000 if prostate_rate else None,
        'lung_implied_cases': lung_rate * population / 100000 if lung_rate else None,
        # Cardiovascular
        'mi_incidence': r.get('MI incidence /100k'),
        'mi_prevalence': r.get('MI prevalence /100k'),
        'heart_care_quality': r.get('Heart care quality (0-100)'),
        # Infectious disease
        'tbe_cases_2025': r.get('TBE cases 2025'),
        'tbe_per100k_2025': r.get('TBE /100k 2025'),
        'ipd_cases_2025': r.get('IPD cases 2025'),
        'ipd_per100k_2025': r.get('IPD /100k 2025'),
        'antibiotic_rx_per1000': r.get('Antibiotic Rx /1000'),
        'mpr_2yo_pct': r.get('MPR 2yo %'),
        'hpv_girls_pct': r.get('HPV girls %'),
        # Pfizer
        'pfizer_total_3yr': f.get('Total Pfizer SEK'),
        'pfizer_per_100k_3yr': f.get('Total Pfizer SEK per 100k'),
        'pfizer_rank': f.get('Pfizer SEK rank'),
        'pfizer_vaccines': f.get('Vaccines SEK'),
        'pfizer_rdim': f.get('RD/IM SEK'),
        'pfizer_oncology': f.get('Oncology SEK'),
        'top_product': t.get('Top Pfizer product'),
        'archetype': a.get('Primary archetype'),
        # Vaccines
        'rsv_share': v.get('rsv_share'),
        'pneu_share': v.get('pneu_share'),
        'tbe_share': v.get('tbe_share'),
        # Per-product 3yr SEK
        'abrysvo': f.get('ABRYSVO SEK (3yr)'),
        'prevenar20': f.get('PREVENAR 20 SEK (3yr)'),
        'fsme_vuxen': f.get('FSME-IMMUN VUXEN SEK (3yr)'),
        'fsme_junior': f.get('FSME-IMMUN JUNIOR SEK (3yr)'),
        'vyndaqel': f.get('VYNDAQEL SEK (3yr)'),
        'vydura': f.get('VYDURA SEK (3yr)'),
        'ibrance': f.get('IBRANCE SEK (3yr)'),
        'tukysa': f.get('TUKYSA SEK (3yr)'),
        'talzenna': f.get('TALZENNA SEK (3yr)'),
        'lorviqua': f.get('LORVIQUA SEK (3yr)'),
        'elrexfio': f.get('ELREXFIO SEK (3yr)'),
        'xtandi': f.get('XTANDI SEK (3yr)'),
    })

# Now also add the equity columns (post_secondary, premature_mortality, etc) from Region master
# Find them dynamically since the headers have Swedish chars
for row in rows:
    region = row['region']
    r = rm[region]
    for k, v in r.items():
        if k:
            ks = str(k)
            if 'Post-secondary' in ks:
                row['post_secondary_pct'] = v
            elif 'Premature mortality' in ks:
                row['premature_mortality_25_64'] = v
            elif 'amenable mortality' in ks.lower():
                row['hc_amenable_mortality'] = v
            elif 'Suicide' in ks:
                row['suicide_25plus'] = v

rows.sort(key=lambda r: r['pfizer_rank'] if r['pfizer_rank'] else 99)

with open('working/data/master/deck_model.json', 'w', encoding='utf-8') as fh:
    json.dump(rows, fh, ensure_ascii=False, indent=2, default=str)

# Verify with Stockholm
sthlm = next(r for r in rows if r['region'] == 'Region Stockholm')
print('Stockholm sample fields:')
for k in ['pop', 'pop_2040', 'pop_65_2040', 'hc_cost_per_capita', 'pharma_per_capita',
          'foreign_born_pct', 'post_secondary_pct', 'premature_mortality_25_64',
          'hc_amenable_mortality', 'suicide_25plus', 'breast_ca_rate', 'mi_incidence',
          'tbe_cases_2025', 'specialists', 'births_2024']:
    print(f'  {k}: {sthlm.get(k)}')

print(f'\n{len(rows)} regions, {len(rows[0])} fields each')

# Check for missing fields by counting None values for each field across regions
print('\nField completeness across 21 regions:')
field_counts = {}
for r in rows:
    for k, v in r.items():
        if v is None or v == 'None':
            field_counts[k] = field_counts.get(k, 0) + 1

for k, count in sorted(field_counts.items(), key=lambda x: -x[1]):
    if count > 0:
        print(f'  {k}: {count}/21 missing')
