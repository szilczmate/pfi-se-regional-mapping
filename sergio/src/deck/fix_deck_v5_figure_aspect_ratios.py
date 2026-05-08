# Author: Sergio Flores - Viti Science - VS-2026-PFI-001
"""Fix stretched figures on slides 4-8 of deck v5 (inherited from v4).

Strategy:
  - Detect each slide's sidebar (any text shape with left >= 9.0in & top > 1.5in)
  - Compute figure-zone: x = 0.5 to (sidebar_left - 0.15) if sidebar exists, else 0.5 to 12.83
  - Vertical zone: y = 1.55 (below subtitle) to 6.80 (above footer/caption)
  - Pick figure size preserving original aspect ratio, fit within zone, centered
"""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
from PIL import Image
import io

ROOT = Path(r"C:/Users/sergi/OneDrive/Desktop/Viti Science/pfi-se-regional-mapping-main")
V5 = ROOT / "delivery" / "08_PFIZER_THURSDAY_DECK_v5.pptx"

SLIDE_W_IN = 13.333
SLIDE_H_IN = 7.5
TOP_BAND = 1.55      # heading + subtitle finishes here
BOTTOM_BAND = 6.80   # footer/caption starts here


def detect_sidebar_left(slide):
    """Find leftmost x of a sidebar text shape (x >= 9.0in & between 1.5 and 7.0 vertically)."""
    sidebar_lefts = []
    for shape in slide.shapes:
        if shape.shape_type == 13:
            continue
        if not shape.has_text_frame or not shape.text_frame.text.strip():
            continue
        x = shape.left / 914400
        y = shape.top / 914400
        h = shape.height / 914400
        if x >= 9.0 and 1.5 <= y <= 7.0:
            sidebar_lefts.append(x)
    return min(sidebar_lefts) if sidebar_lefts else None


def fix_picture(pic, sidebar_left):
    blob = pic.image.blob
    img = Image.open(io.BytesIO(blob))
    orig_ratio = img.size[0] / img.size[1]

    cur_w = pic.width / 914400
    cur_h = pic.height / 914400

    # Determine zone
    zone_left = 0.5
    zone_right = (sidebar_left - 0.15) if sidebar_left else (SLIDE_W_IN - 0.5)
    zone_top = TOP_BAND
    zone_bottom = BOTTOM_BAND
    zone_w = zone_right - zone_left
    zone_h = zone_bottom - zone_top

    # Pick max figure size preserving aspect ratio
    new_w = zone_w
    new_h = new_w / orig_ratio
    if new_h > zone_h:
        new_h = zone_h
        new_w = new_h * orig_ratio

    # Center within the zone
    new_left = zone_left + (zone_w - new_w) / 2
    new_top = zone_top + (zone_h - new_h) / 2

    pic.left = Inches(new_left)
    pic.top = Inches(new_top)
    pic.width = Inches(new_w)
    pic.height = Inches(new_h)

    return cur_w, cur_h, new_w, new_h, new_left, new_top, zone_left, zone_right


def main():
    prs = Presentation(str(V5))
    print(f'Loaded {len(prs.slides)} slides\n')

    for idx in [4, 5, 6, 7, 8]:
        sld = prs.slides[idx - 1]
        sidebar_left = detect_sidebar_left(sld)
        for shape in sld.shapes:
            if shape.shape_type == 13:
                cw, ch, nw, nh, nl, nt, zl, zr = fix_picture(shape, sidebar_left)
                sb_str = f'sidebar @ x={sidebar_left:.2f}' if sidebar_left else 'no sidebar'
                print(f'  Slide {idx} ({sb_str}, zone x={zl:.2f}-{zr:.2f}): '
                      f'{cw:.2f}x{ch:.2f} → {nw:.2f}x{nh:.2f} at ({nl:.2f},{nt:.2f})')
                break

    prs.save(str(V5))
    print(f'\nSaved {V5}')


if __name__ == '__main__':
    main()
